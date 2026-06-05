"""Generate drone_<D>.xlsx for D in 0..33 from vut.lp, mirroring drone_31.xlsx.

Columns:
  A: next(D, X, V)
  B: arrival(D, V, X, T)              (blank for X=0)
  C: start(D, V, X, T)
  D: FT = flight_time((V_{X-1}, V_X)) (blank for X=0 -> "FT" header; blank if 0)
  E: CT = charge_time((V_{X-1}, V_X)) (blank for X=0 -> "CT" header; blank if 0)

After two blank rows, all in(...) facts are listed in column A, with the
current drone's own in/1 records first.
"""

import re
from pathlib import Path

import openpyxl

HERE = Path(__file__).parent
VUT = HERE / "vut.lp"
NETWORK = HERE.parent / "instances" / "network_NY.lp"
MAX_SEG = 24
N_AGENTS = 34


def parse_network():
    ft, ct = {}, {}
    pat_ft = re.compile(r"flight_time\(\(\s*(\w+)\s*,\s*(\w+)\s*\)\s*,\s*(\d+)\s*\)")
    pat_ct = re.compile(r"charge_time\(\(\s*(\w+)\s*,\s*(\w+)\s*\)\s*,\s*(\d+)\s*\)")
    text = NETWORK.read_text()
    for v1, v2, t in pat_ft.findall(text):
        ft[(v1, v2)] = int(t)
    for v1, v2, t in pat_ct.findall(text):
        ct[(v1, v2)] = int(t)
    return ft, ct


def parse_vut():
    text = VUT.read_text()
    nxt = {}      # (D, X) -> V
    arr = {}      # (D, X) -> (V, T)
    sta = {}      # (D, X) -> (V, T)
    pat_next = re.compile(r"^next\((\d+),(\d+),(\w+)\)$", re.M)
    pat_arr = re.compile(r"^arrival\((\d+),(\w+),(\d+),(\d+)\)$", re.M)
    pat_sta = re.compile(r"^start\((\d+),(\w+),(\d+),(\d+)\)$", re.M)
    pat_in = re.compile(r"^in\(.+?\)$", re.M)
    for d, x, v in pat_next.findall(text):
        nxt[(int(d), int(x))] = v
    for d, v, x, t in pat_arr.findall(text):
        arr[(int(d), int(x))] = (v, int(t))
    for d, v, x, t in pat_sta.findall(text):
        sta[(int(d), int(x))] = (v, int(t))
    ins = pat_in.findall(text)
    return nxt, arr, sta, ins


def own_in_records(ins, drone):
    own, other = [], []
    pat = re.compile(r"^in\((?:arrival|start)\((\d+),")
    for s in ins:
        m = pat.match(s)
        if m and int(m.group(1)) == drone:
            own.append(s)
        else:
            other.append(s)
    return own + other


def write_drone(drone, nxt, arr, sta, ins, ft, ct):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    for x in range(MAX_SEG + 1):
        v = nxt.get((drone, x))
        if v is None:
            continue
        row = x + 1
        ws.cell(row=row, column=1, value=f"next({drone},{x},{v})")
        if x == 0:
            sv, st = sta.get((drone, 0), (v, 0))
            ws.cell(row=row, column=3, value=f"start({drone},{sv},0,{st})")
            ws.cell(row=row, column=4, value="FT")
            ws.cell(row=row, column=5, value="CT")
        else:
            av, at = arr[(drone, x)]
            sv, st = sta[(drone, x)]
            ws.cell(row=row, column=2, value=f"arrival({drone},{av},{x},{at})")
            ws.cell(row=row, column=3, value=f"start({drone},{sv},{x},{st})")
            prev = nxt[(drone, x - 1)]
            ft_val = ft.get((prev, v), 0)
            ct_val = ct.get((prev, v), 0)
            if ft_val:
                ws.cell(row=row, column=4, value=ft_val)
            if ct_val:
                ws.cell(row=row, column=5, value=ct_val)

    in_rows = own_in_records(ins, drone)
    start_row = MAX_SEG + 1 + 3
    for i, s in enumerate(in_rows):
        ws.cell(row=start_row + i, column=1, value=s)

    out = HERE / f"drone_{drone}.xlsx"
    wb.save(out)
    return out


def main():
    ft, ct = parse_network()
    nxt, arr, sta, ins = parse_vut()
    for d in range(N_AGENTS):
        path = write_drone(d, nxt, arr, sta, ins, ft, ct)
        print(f"wrote {path.name}")


if __name__ == "__main__":
    main()
